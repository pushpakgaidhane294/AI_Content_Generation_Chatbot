with open("app/routes/chat.py", "r", encoding="utf-8") as f:
    py = f.read()

import re

old_logic = """    doc = Document()
    for line in req.text.split("\\n"):
        doc.add_paragraph(line)"""

new_logic = """    import re
    doc = Document()
    
    def _parse_inline(p, text):
        pattern = re.compile(r'(\*\*(.*?)\*\*|\*(.*?)\*|`(.*?)`|\[(.*?)\]\((.*?)\))')
        last_idx = 0
        for match in pattern.finditer(text):
            if match.start() > last_idx:
                p.add_run(text[last_idx:match.start()])
            full_match = match.group(0)
            if full_match.startswith('**'):
                run = p.add_run(match.group(2))
                run.bold = True
            elif full_match.startswith('*'):
                run = p.add_run(match.group(3))
                run.italic = True
            elif full_match.startswith('`'):
                run = p.add_run(match.group(4))
                try:
                    from docx.shared import Pt
                    run.font.name = 'Courier New'
                except:
                    pass
            elif full_match.startswith('['):
                run = p.add_run(match.group(5))
                run.underline = True
                try:
                    from docx.shared import RGBColor
                    run.font.color.rgb = RGBColor(5, 99, 193)
                except:
                    pass
            last_idx = match.end()
        if last_idx < len(text):
            p.add_run(text[last_idx:])

    lines = req.text.split('\\n')
    i = 0
    in_code_block = False
    code_block_text = []
    
    while i < len(lines):
        line = lines[i].strip()
        
        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph('\\n'.join(code_block_text))
                p.style = 'No Spacing' if 'No Spacing' in [s.name for s in doc.styles] else 'Normal'
                code_block_text = []
            else:
                in_code_block = True
            i += 1
            continue
            
        if in_code_block:
            code_block_text.append(lines[i])
            i += 1
            continue
            
        if not line:
            i += 1
            continue
            
        if line == '---' or line == '***' or line == '___':
            doc.add_paragraph('_' * 40)
            i += 1
            continue
            
        m = re.match(r'^(#{1,6})\s+(.*)', line)
        if m:
            level = len(m.group(1))
            heading_text = m.group(2)
            p = doc.add_heading(level=level)
            _parse_inline(p, heading_text)
            i += 1
            continue
            
        m = re.match(r'^[\-\*]\s+(.*)', line)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            _parse_inline(p, m.group(1))
            i += 1
            continue
            
        m = re.match(r'^\d+\.\s+(.*)', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            _parse_inline(p, m.group(1))
            i += 1
            continue
            
        if line.startswith('|') and i + 1 < len(lines) and lines[i+1].strip().startswith('|'):
            # Check if second row is a separator row
            if '---' in lines[i+1]:
                headers = [h for h in line.split('|')[1:-1]]
                table_lines = []
                i += 2
                while i < len(lines) and lines[i].strip().startswith('|'):
                    table_lines.append(lines[i].strip())
                    i += 1
                    
                if headers:
                    table = doc.add_table(rows=1, cols=len(headers))
                    table.style = 'Table Grid'
                    hdr_cells = table.rows[0].cells
                    for j, h in enumerate(headers):
                        if j < len(hdr_cells):
                            _parse_inline(hdr_cells[j].paragraphs[0], h.strip())
                            if hdr_cells[j].paragraphs[0].runs:
                                hdr_cells[j].paragraphs[0].runs[0].bold = True
                            
                    for tline in table_lines:
                        row_cells = table.add_row().cells
                        cells = [c for c in tline.split('|')[1:-1]]
                        for j, c in enumerate(cells):
                            if j < len(row_cells):
                                _parse_inline(row_cells[j].paragraphs[0], c.strip())
                continue

        p = doc.add_paragraph()
        _parse_inline(p, line)
        i += 1"""

py = py.replace(old_logic, new_logic)

with open("app/routes/chat.py", "w", encoding="utf-8") as f:
    f.write(py)
