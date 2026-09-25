import re

def parse_markdown_to_docx(doc, text):
    lines = text.split('\n')
    i = 0
    in_code_block = False
    code_block_text = []
    
    while i < len(lines):
        line = lines[i].strip()
        
        # Code Blocks
        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph('\n'.join(code_block_text))
                p.style = 'No Spacing' # Or 'Macro Text'
                code_block_text = []
            else:
                in_code_block = True
            i += 1
            continue
            
        if in_code_block:
            code_block_text.append(lines[i]) # Keep original indentation
            i += 1
            continue
            
        if not line:
            i += 1
            continue
            
        # Horizontal Rule
        if line == '---' or line == '***' or line == '___':
            doc.add_paragraph('_' * 40)
            i += 1
            continue
            
        # Headings
        m = re.match(r'^(#{1,6})\s+(.*)', line)
        if m:
            level = len(m.group(1))
            heading_text = m.group(2)
            p = doc.add_heading(level=level)
            _parse_inline(p, heading_text)
            i += 1
            continue
            
        # Bullet Lists
        m = re.match(r'^[\-\*]\s+(.*)', line)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            _parse_inline(p, m.group(1))
            i += 1
            continue
            
        # Numbered Lists
        m = re.match(r'^\d+\.\s+(.*)', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            _parse_inline(p, m.group(1))
            i += 1
            continue
            
        # Tables
        if line.startswith('|') and i + 1 < len(lines) and lines[i+1].strip().startswith('|---'):
            # Parse table headers
            headers = [cell.strip() for cell in line.split('|') if cell.strip() or not cell == '']
            headers = [h for h in line.split('|')[1:-1]]
            
            # Find all table rows
            table_lines = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
                
            if headers:
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                for j, h in enumerate(headers):
                    if j < len(hdr_cells):
                        _parse_inline(hdr_cells[j].paragraphs[0], h.strip())
                        hdr_cells[j].paragraphs[0].runs[0].bold = True if hdr_cells[j].paragraphs[0].runs else False
                        
                for tline in table_lines:
                    row_cells = table.add_row().cells
                    cells = [c for c in tline.split('|')[1:-1]]
                    for j, c in enumerate(cells):
                        if j < len(row_cells):
                            _parse_inline(row_cells[j].paragraphs[0], c.strip())
            continue

        # Normal Paragraph
        p = doc.add_paragraph()
        _parse_inline(p, line)
        i += 1

def _parse_inline(p, text):
    # Regex for links, bold, italic, code
    # This is a basic tokenizer
    pattern = re.compile(r'(\*\*(.*?)\*\*|\*(.*?)\*|`(.*?)`|\[(.*?)\]\((.*?)\))')
    
    last_idx = 0
    for match in pattern.finditer(text):
        # Add text before match
        if match.start() > last_idx:
            p.add_run(text[last_idx:match.start()])
            
        full_match = match.group(0)
        
        if full_match.startswith('**'):
            run = p.add_run(match.group(2))
            run.bold = True
        elif full_match.startswith('*'):
            run = p.add_run(match.group(3))
            run.italic = True
        elif full_match.startswith('`'):
            run = p.add_run(match.group(4))
            run.font.name = 'Courier New'
        elif full_match.startswith('['):
            # link
            link_text = match.group(5)
            # URL = match.group(6)
            # python-docx doesn't do hyperlinks easily without complex xml. Just adding text.
            run = p.add_run(link_text)
            run.underline = True
            run.font.color.rgb = __import__('docx').shared.RGBColor(5, 99, 193)
            
        last_idx = match.end()
        
    if last_idx < len(text):
        p.add_run(text[last_idx:])

# Test code
if __name__ == '__main__':
    from docx import Document
    doc = Document()
    test_md = """# AI Content Generation Chatbot
## Features
This application uses **Generative AI** and *Prompt Engineering*.
### Main Features
- Email Generation
- Report Generation
- Technical Explanation
### Technologies
| Technology | Purpose |
|------------|---------|
| FastAPI    | Backend |
| Groq       | AI Generation |
| SQLite     | Database |
### Example Code
```python
print("Hello World")
```
Visit [Google](https://www.google.com/)"""
    parse_markdown_to_docx(doc, test_md)
    doc.save("test.docx")
    print("Done")
